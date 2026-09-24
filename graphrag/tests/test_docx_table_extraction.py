import os
import tempfile
import unittest

import docx

from common.utils.text_extractors import (
    ExtractionError,
    _docx_to_markdown,
    extract_text_from_file,
    get_supported_extensions,
)


def _saved(build):
    """Build a .docx with ``build(document)`` and return its path."""
    handle = tempfile.NamedTemporaryFile(suffix=".docx", delete=False)
    handle.close()
    document = docx.Document()
    build(document)
    document.save(handle.name)
    return handle.name


class TestDocxTablesSurvive(unittest.TestCase):
    """Tables in a .docx must reach the extracted text.

    Extraction previously joined ``doc.paragraphs`` only (GML-2196). That
    yields just the direct children of the document body, so every table was
    dropped — silently, because the prose around it still came through.
    """

    def setUp(self):
        self.paths = []

    def tearDown(self):
        for path in self.paths:
            try:
                os.unlink(path)
            except OSError:
                pass

    def _build(self, build):
        path = _saved(build)
        self.paths.append(path)
        return path

    @staticmethod
    def _thresholds(document):
        document.add_paragraph("Approval thresholds are defined below.")
        table = document.add_table(rows=3, cols=2)
        rows = [("Tier", "Limit"), ("Manager", "$10,000"), ("VP", "$250,000")]
        for index, (left, right) in enumerate(rows):
            table.rows[index].cells[0].text = left
            table.rows[index].cells[1].text = right
        document.add_paragraph("Requests above the VP limit escalate to the CFO.")

    def test_cell_values_are_extracted(self):
        """The figures a reader would ask about are present."""
        text = extract_text_from_file(self._build(self._thresholds))
        for value in ("Tier", "Limit", "Manager", "$10,000", "VP", "$250,000"):
            self.assertIn(value, text)

    def test_table_stays_between_its_paragraphs(self):
        """A table keeps its position, so the prose introducing it still refers
        to something."""
        text = extract_text_from_file(self._build(self._thresholds))
        intro = text.index("Approval thresholds")
        table = text.index("$250,000")
        outro = text.index("escalate to the CFO")
        self.assertLess(intro, table)
        self.assertLess(table, outro)

    def test_rendered_as_a_markdown_table(self):
        """Tabular data is rendered the way the spreadsheet branch renders it."""
        text = extract_text_from_file(self._build(self._thresholds))
        self.assertIn("| Tier | Limit |", text)
        self.assertIn("| VP | $250,000 |", text)

    def test_pipe_and_newline_in_a_cell_do_not_break_the_row(self):
        """A cell containing a pipe or a line break must not split the row."""

        def build(document):
            table = document.add_table(rows=2, cols=2)
            table.rows[0].cells[0].text = "Region"
            table.rows[0].cells[1].text = "Rule"
            table.rows[1].cells[0].text = "EU|UK"
            cell = table.rows[1].cells[1]
            cell.text = "line one"
            cell.add_paragraph("line two")

        text = extract_text_from_file(self._build(build))
        rows = [line for line in text.splitlines() if line.startswith("|")]
        self.assertEqual(len(rows), 3, f"expected header, rule and one row: {rows}")
        self.assertIn(r"EU\|UK", text)
        self.assertIn("line one line two", text)

    def test_nested_table_contents_are_reached(self):
        """A table inside a cell is flattened rather than dropped."""

        def build(document):
            outer = document.add_table(rows=1, cols=1)
            inner = outer.rows[0].cells[0].add_table(rows=1, cols=2)
            inner.rows[0].cells[0].text = "NESTED_KEY"
            inner.rows[0].cells[1].text = "NESTED_VAL"

        text = extract_text_from_file(self._build(build))
        self.assertIn("NESTED_KEY", text)
        self.assertIn("NESTED_VAL", text)

    def test_merged_cells_do_not_raise(self):
        """Merged cells repeat rather than producing a ragged row."""

        def build(document):
            table = document.add_table(rows=2, cols=3)
            table.rows[0].cells[0].merge(table.rows[0].cells[1])
            table.rows[0].cells[0].text = "Merged"
            table.rows[0].cells[2].text = "C"
            for index, value in enumerate(["x", "y", "z"]):
                table.rows[1].cells[index].text = value

        text = extract_text_from_file(self._build(build))
        self.assertIn("Merged", text)
        self.assertIn("| x | y | z |", text)

    def test_document_without_tables_is_unchanged(self):
        """Prose-only documents keep their previous shape."""

        def build(document):
            document.add_paragraph("First.")
            document.add_paragraph("Second.")

        self.assertEqual(
            _docx_to_markdown(docx.Document(self._build(build))),
            "First.\n\nSecond.",
        )


class TestUnreadableFilesRaise(unittest.TestCase):
    """Formats that cannot be read must raise, not become a document."""

    def _write(self, suffix, payload):
        handle = tempfile.NamedTemporaryFile(suffix=suffix, delete=False)
        handle.write(payload)
        handle.close()
        self.addCleanup(lambda: os.path.exists(handle.name) and os.unlink(handle.name))
        return handle.name

    def test_legacy_doc_raises(self):
        """.doc is the OLE format python-docx cannot read."""
        path = self._write(".doc", b"\xd0\xcf\x11\xe0 legacy")
        with self.assertRaises(ExtractionError) as caught:
            extract_text_from_file(path)
        self.assertIn(".docx", str(caught.exception))

    def test_legacy_doc_is_not_advertised_as_supported(self):
        self.assertNotIn(".doc", get_supported_extensions())
        self.assertIn(".docx", get_supported_extensions())

    def test_unsupported_extension_raises(self):
        path = self._write(".rtf", b"{}")
        with self.assertRaises(ExtractionError):
            extract_text_from_file(path)

    def test_corrupt_docx_raises_without_leaking_parser_detail(self):
        """The caller-facing message must not carry the parser's wording."""
        path = self._write(".docx", b"not a zip at all")
        with self.assertRaises(ExtractionError) as caught:
            extract_text_from_file(path)
        message = str(caught.exception)
        self.assertNotIn("zip", message.lower())
        self.assertIn("corrupt", message.lower())


if __name__ == "__main__":
    unittest.main()
